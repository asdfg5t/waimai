#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

import sys
import re
from pathlib import Path
from typing import List

try:
	from docx import Document
	from docx.shared import Pt
	from docx.oxml.shared import OxmlElement, qn
	from docx.enum.text import WD_ALIGN_PARAGRAPH
except Exception as exc:  # pragma: no cover
	print("python-docx 未安装，请先运行: pip install python-docx", file=sys.stderr)
	raise


HEADING_STYLE_BY_LEVEL = {
	1: "Heading 1",
	2: "Heading 2",
	3: "Heading 3",
	4: "Heading 4",
	5: "Heading 5",
	6: "Heading 6",
}

BULLET_STYLE = "List Bullet"
NUMBER_STYLE = "List Number"
QUOTE_STYLE = "Intense Quote"


def add_horizontal_rule(paragraph):
	"""Add a bottom border to emulate a horizontal rule."""
	p = paragraph._p
	p_pr = p.get_or_add_pPr()
	p_borders = OxmlElement('w:pBdr')
	p_bdr = OxmlElement('w:bottom')
	p_bdr.set(qn('w:val'), 'single')
	p_bdr.set(qn('w:sz'), '6')  # size
	p_bdr.set(qn('w:space'), '1')
	p_bdr.set(qn('w:color'), 'auto')
	p_borders.append(p_bdr)
	p_pr.append(p_borders)


def parse_markdown_table_block(table_lines: List[str]) -> List[List[str]]:
	"""Convert a list of markdown table lines into a 2D array (rows x cols)."""
	rows: List[List[str]] = []
	for idx, raw in enumerate(table_lines):
		line = raw.strip()
		if not line.startswith('|'):
			continue
		# split and drop empty first/last that come from leading/trailing '|'
		cells = [c.strip() for c in line.split('|')]
		if cells and cells[0] == '':
			cells = cells[1:]
		if cells and cells[-1] == '':
			cells = cells[:-1]
		# skip separator row like | --- | --- |
		if all(re.match(r"^:?-{3,}:?$", c) for c in cells):
			continue
		rows.append(cells)
	return rows


def write_table(document: Document, rows: List[List[str]]):
	if not rows:
		return
	col_count = max(len(r) for r in rows)
	table = document.add_table(rows=len(rows), cols=col_count)
	table.style = 'Table Grid'
	for r_idx, r in enumerate(rows):
		for c_idx in range(col_count):
			text = r[c_idx] if c_idx < len(r) else ''
			cell = table.cell(r_idx, c_idx)
			cell.text = text


def convert_markdown_to_docx(markdown_path: Path, docx_path: Path):
	text = markdown_path.read_text(encoding='utf-8')
	document = Document()

	# Basic normal style font size tweak for readability
	style = document.styles['Normal']
	style.font.name = '宋体'
	style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')  # type: ignore[attr-defined]
	style.font.size = Pt(11)

	lines = text.splitlines()
	in_code_fence = False
	code_fence_delim = ''
	table_buffer: List[str] = []

	def flush_table_if_any():
		nonlocal table_buffer
		if table_buffer:
			rows = parse_markdown_table_block(table_buffer)
			write_table(document, rows)
			table_buffer = []

	for raw_line in lines:
		line = raw_line.rstrip('\n')

		# Handle fenced code blocks pass-through as preformatted paragraphs
		m_fence = re.match(r"^(```+)(.*)$", line)
		if m_fence:
			flush_table_if_any()
			delim = m_fence.group(1)
			if not in_code_fence:
				in_code_fence = True
				code_fence_delim = delim
				p = document.add_paragraph()
				p.style = style
				continue
			else:
				if delim == code_fence_delim:
					in_code_fence = False
					code_fence_delim = ''
					continue

		if in_code_fence:
			p = document.add_paragraph(line)
			p.style = 'No Spacing'
			continue

		# Accumulate markdown table lines
		if line.strip().startswith('|'):
			table_buffer.append(line)
			continue
		else:
			flush_table_if_any()

		# Horizontal rule
		if re.match(r"^\s*[-*_]{3,}\s*$", line):
			p = document.add_paragraph()
			add_horizontal_rule(p)
			continue

		# Headings
		m = re.match(r"^(#{1,6})\s+(.*)$", line)
		if m:
			level = len(m.group(1))
			text_content = m.group(2).strip()
			p = document.add_paragraph(text_content)
			p.style = HEADING_STYLE_BY_LEVEL.get(level, 'Heading 1')
			continue

		# Blockquote
		m = re.match(r"^>\s?(.*)$", line)
		if m:
			p = document.add_paragraph(m.group(1).strip())
			try:
				p.style = QUOTE_STYLE
			except Exception:
				pass
			continue

		# Unordered list with checkbox
		m = re.match(r"^\s*[-*]\s+\[( |x|X)\]\s+(.*)$", line)
		if m:
			checked = m.group(1).lower() == 'x'
			content = m.group(2)
			prefix = '■ ' if checked else '□ '
			p = document.add_paragraph(f"{prefix}{content}")
			p.style = BULLET_STYLE
			continue

		# Unordered list
		m = re.match(r"^\s*[-*]\s+(.*)$", line)
		if m:
			p = document.add_paragraph(m.group(1).strip())
			p.style = BULLET_STYLE
			continue

		# Ordered list
		m = re.match(r"^\s*\d+\.\s+(.*)$", line)
		if m:
			p = document.add_paragraph(m.group(1).strip())
			p.style = NUMBER_STYLE
			continue

		# Empty line -> add spacing
		if line.strip() == "":
			document.add_paragraph("")
			continue

		# Normal paragraph
		p = document.add_paragraph(line)
		p.alignment = WD_ALIGN_PARAGRAPH.LEFT

	# flush any trailing table
	flush_table_if_any()

	docx_path.parent.mkdir(parents=True, exist_ok=True)
	document.save(str(docx_path))


def main():
	if len(sys.argv) != 3:
		print("用法: md_to_docx.py 源.md 目标.docx", file=sys.stderr)
		sys.exit(2)
	md = Path(sys.argv[1]).expanduser().resolve()
	dx = Path(sys.argv[2]).expanduser().resolve()
	if not md.exists():
		print(f"找不到源文件: {md}", file=sys.stderr)
		sys.exit(1)
	convert_markdown_to_docx(md, dx)
	print(f"已生成: {dx}")


if __name__ == "__main__":
	main()